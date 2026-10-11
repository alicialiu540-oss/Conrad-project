#!/usr/bin/env python3
"""Process stretching videos with YOLO Pose and update the dataset workbook."""

from __future__ import annotations

import argparse
import csv
import hashlib
import json
import os
import re
import sys
import tempfile
from copy import copy
from datetime import datetime
from pathlib import Path
from typing import Any

import cv2
import numpy as np
import openpyxl
import torch
from ultralytics import YOLO, __version__ as ultralytics_version


ROOT = Path(__file__).resolve().parent.parent
DEFAULT_INBOX = ROOT / "data" / "inbox"
DEFAULT_MODEL = ROOT / "models" / "yolo26x-pose.pt"
DEFAULT_WORKBOOK = ROOT / "outputs" / "conrad-pose-dataset" / "Conrad_Pose_Dataset_Tracker.xlsx"
KEYPOINT_DIR = ROOT / "data" / "processed" / "keypoints"
ANNOTATED_DIR = ROOT / "data" / "processed" / "annotated_videos"
MANIFEST_PATH = ROOT / "data" / "processed" / "processed_videos.json"
SUPPORTED_VIDEO_EXTENSIONS = {".mp4", ".mov", ".m4v", ".avi", ".mkv", ".webm"}
KEYPOINT_NAMES = [
    "nose", "left_eye", "right_eye", "left_ear", "right_ear",
    "left_shoulder", "right_shoulder", "left_elbow", "right_elbow",
    "left_wrist", "right_wrist", "left_hip", "right_hip",
    "left_knee", "right_knee", "left_ankle", "right_ankle",
]
# A side-view exercise can be fully usable even when the far-side eye, ear, arm,
# or leg is hidden. Count the body as visible when at least one keypoint from
# every major left/right joint pair is confidently detected.
BODY_JOINT_GROUPS = [
    (5, 6),    # shoulders
    (7, 8),    # elbows
    (9, 10),   # wrists
    (11, 12),  # hips
    (13, 14),  # knees
    (15, 16),  # ankles
]
CAMERA_VIEWS = {"front", "side_left", "side_right", "back", "diagonal"}


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description="Run YOLO Pose on videos and automatically update the Conrad dataset workbook."
    )
    parser.add_argument("inputs", nargs="*", type=Path, help="Video files or folders. Defaults to data/inbox.")
    parser.add_argument("--model", type=Path, default=DEFAULT_MODEL)
    parser.add_argument("--workbook", type=Path, default=DEFAULT_WORKBOOK)
    parser.add_argument("--participant-id", help="Anonymous ID, for example P001. Applied to all input videos.")
    parser.add_argument("--exercise", help="Exercise label. Applied to all input videos.")
    parser.add_argument("--camera-view", choices=sorted(CAMERA_VIEWS), help="Camera view for all input videos.")
    parser.add_argument(
        "--consent",
        choices=["yes", "no", "reference_only"],
        help="Whether the videos may be used, or analyzed locally as reference-only material.",
    )
    parser.add_argument("--device", default="auto", help="auto, mps, cpu, or a CUDA device such as 0.")
    parser.add_argument("--imgsz", type=int, default=640, help="YOLO input size. Default: 640.")
    parser.add_argument("--confidence", type=float, default=0.25, help="Person detection threshold.")
    parser.add_argument("--keypoint-confidence", type=float, default=0.50, help="Visible-keypoint threshold.")
    parser.add_argument("--non-interactive", action="store_true", help="Use placeholders instead of asking questions.")
    return parser.parse_args()


def choose_device(requested: str) -> str:
    if requested != "auto":
        return requested
    if torch.cuda.is_available():
        return "0"
    if torch.backends.mps.is_available():
        return "mps"
    return "cpu"


def discover_videos(inputs: list[Path]) -> list[Path]:
    sources = inputs or [DEFAULT_INBOX]
    found: list[Path] = []
    for source in sources:
        source = source.expanduser().resolve()
        if source.is_file() and source.suffix.lower() in SUPPORTED_VIDEO_EXTENSIONS:
            found.append(source)
        elif source.is_dir():
            found.extend(
                path.resolve()
                for path in sorted(source.iterdir())
                if path.is_file() and path.suffix.lower() in SUPPORTED_VIDEO_EXTENSIONS
            )
        elif source.exists():
            print(f"Skipping unsupported input: {source}")
        else:
            print(f"Input does not exist: {source}")
    return sorted(set(found))


def sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def load_manifest() -> dict[str, Any]:
    if not MANIFEST_PATH.exists():
        return {"version": 1, "processed": {}}
    try:
        data = json.loads(MANIFEST_PATH.read_text(encoding="utf-8"))
        if not isinstance(data.get("processed"), dict):
            raise ValueError("processed must be an object")
        return data
    except (json.JSONDecodeError, ValueError) as exc:
        raise RuntimeError(f"Manifest is damaged: {MANIFEST_PATH}: {exc}") from exc


def save_manifest(manifest: dict[str, Any]) -> None:
    MANIFEST_PATH.parent.mkdir(parents=True, exist_ok=True)
    temp_path = MANIFEST_PATH.with_suffix(".tmp")
    temp_path.write_text(json.dumps(manifest, indent=2), encoding="utf-8")
    os.replace(temp_path, MANIFEST_PATH)


def clean_label(value: str) -> str:
    value = value.strip().lower()
    value = re.sub(r"[^a-z0-9]+", "_", value)
    return value.strip("_") or "unlabeled"


def parse_filename_metadata(path: Path) -> tuple[str | None, str | None, str | None]:
    # Optional naming format: P001__back_extension__side_left.mp4
    parts = path.stem.split("__")
    if len(parts) < 3:
        return None, None, None
    participant = parts[0].strip() or None
    exercise = clean_label(parts[1])
    camera = clean_label(parts[2])
    if camera not in CAMERA_VIEWS:
        camera = None
    return participant, exercise, camera


def ask(prompt: str, default: str | None = None) -> str:
    suffix = f" [{default}]" if default else ""
    response = input(f"{prompt}{suffix}: ").strip()
    return response or (default or "")


def collect_metadata(path: Path, args: argparse.Namespace) -> dict[str, str]:
    parsed_participant, parsed_exercise, parsed_camera = parse_filename_metadata(path)
    if args.non_interactive:
        return {
            "participant_id": args.participant_id or parsed_participant or "P_UNASSIGNED",
            "exercise_name": clean_label(args.exercise or parsed_exercise or "unlabeled"),
            "camera_view": args.camera_view or parsed_camera or "diagonal",
            "consent": args.consent or "no",
        }

    print(f"\nInformation for {path.name}")
    participant = args.participant_id or parsed_participant or ask("Anonymous participant ID", "P001")
    exercise = args.exercise or parsed_exercise or ask("Exercise name", "unlabeled")
    camera_default = args.camera_view or parsed_camera or "side_left"
    camera = camera_default
    if not args.camera_view and not parsed_camera:
        while True:
            camera = clean_label(ask("Camera view (front/side_left/side_right/back/diagonal)", camera_default))
            if camera in CAMERA_VIEWS:
                break
            print("Please use one of the listed camera views.")
    consent = args.consent
    if consent is None:
        while True:
            consent = ask(
                "Usage status (yes/no/reference_only)", "no"
            ).lower()
            if consent in {"yes", "no", "reference_only"}:
                break
            print("Please enter yes or no.")
    return {
        "participant_id": participant.strip() or "P_UNASSIGNED",
        "exercise_name": clean_label(exercise),
        "camera_view": camera,
        "consent": consent,
    }


def video_properties(path: Path) -> dict[str, Any]:
    capture = cv2.VideoCapture(str(path))
    if not capture.isOpened():
        raise RuntimeError(f"Cannot open video: {path}")
    fps = float(capture.get(cv2.CAP_PROP_FPS) or 0.0)
    frame_count = int(capture.get(cv2.CAP_PROP_FRAME_COUNT) or 0)
    width = int(capture.get(cv2.CAP_PROP_FRAME_WIDTH) or 0)
    height = int(capture.get(cv2.CAP_PROP_FRAME_HEIGHT) or 0)
    capture.release()
    duration = frame_count / fps if fps > 0 else 0.0
    return {"fps": fps, "frame_count": frame_count, "width": width, "height": height, "duration": duration}


def keypoint_header() -> list[str]:
    header = [
        "frame_index", "timestamp_sec", "detected", "person_count",
        "detection_confidence", "bbox_x1", "bbox_y1", "bbox_x2", "bbox_y2",
    ]
    for name in KEYPOINT_NAMES:
        header.extend([f"{name}_x", f"{name}_y", f"{name}_confidence"])
    return header


def safe_float(value: Any) -> float:
    return float(value) if value is not None else 0.0


def process_video(
    model: YOLO,
    path: Path,
    output_stem: str,
    device: str,
    imgsz: int,
    confidence: float,
    keypoint_confidence: float,
) -> dict[str, Any]:
    properties = video_properties(path)
    fps = properties["fps"] or 30.0
    KEYPOINT_DIR.mkdir(parents=True, exist_ok=True)
    ANNOTATED_DIR.mkdir(parents=True, exist_ok=True)
    csv_path = KEYPOINT_DIR / f"{output_stem}_keypoints.csv"
    annotated_path = ANNOTATED_DIR / f"{output_stem}_pose.mp4"

    writer = None
    total_frames = 0
    detected_frames = 0
    full_body_frames = 0
    person_counts: list[int] = []
    detection_confidences: list[float] = []
    keypoint_confidences: list[float] = []

    with csv_path.open("w", newline="", encoding="utf-8") as csv_file:
        csv_writer = csv.writer(csv_file)
        csv_writer.writerow(keypoint_header())
        results = model.predict(
            source=str(path), stream=True, verbose=False, save=False,
            imgsz=imgsz, conf=confidence, device=device,
        )
        for frame_index, result in enumerate(results):
            plotted = result.plot()
            if writer is None:
                height, width = plotted.shape[:2]
                writer = cv2.VideoWriter(
                    str(annotated_path), cv2.VideoWriter_fourcc(*"mp4v"), fps, (width, height)
                )
                if not writer.isOpened():
                    raise RuntimeError(f"Cannot create annotated video: {annotated_path}")
            writer.write(plotted)
            total_frames += 1

            count = 0 if result.keypoints is None else len(result.keypoints)
            person_counts.append(count)
            base = [frame_index, frame_index / fps, 0, count, "", "", "", "", ""]
            if count == 0 or result.boxes is None:
                csv_writer.writerow(base + [""] * (len(KEYPOINT_NAMES) * 3))
                continue

            boxes = result.boxes.xyxy.detach().cpu().numpy()
            box_conf = result.boxes.conf.detach().cpu().numpy()
            areas = (boxes[:, 2] - boxes[:, 0]) * (boxes[:, 3] - boxes[:, 1])
            primary_index = int(np.argmax(areas))
            points = result.keypoints.data[primary_index].detach().cpu().numpy()
            detected_frames += 1
            detection_confidences.append(safe_float(box_conf[primary_index]))
            point_conf = points[:, 2] if points.shape[1] >= 3 else np.ones(len(points))
            keypoint_confidences.extend(float(value) for value in point_conf)
            if len(point_conf) >= len(KEYPOINT_NAMES) and all(
                any(point_conf[index] >= keypoint_confidence for index in group)
                for group in BODY_JOINT_GROUPS
            ):
                full_body_frames += 1

            box = boxes[primary_index]
            row: list[Any] = [
                frame_index, frame_index / fps, 1, count, safe_float(box_conf[primary_index]),
                safe_float(box[0]), safe_float(box[1]), safe_float(box[2]), safe_float(box[3]),
            ]
            for index in range(len(KEYPOINT_NAMES)):
                if index < len(points):
                    row.extend([
                        safe_float(points[index, 0]), safe_float(points[index, 1]), safe_float(point_conf[index])
                    ])
                else:
                    row.extend(["", "", ""])
            csv_writer.writerow(row)

    if writer is not None:
        writer.release()
    else:
        raise RuntimeError(f"No frames were read from {path}")

    detection_coverage = detected_frames / total_frames if total_frames else 0.0
    full_body_coverage = full_body_frames / total_frames if total_frames else 0.0
    if full_body_coverage >= 0.80:
        whole_body_visible = "yes"
    elif detection_coverage >= 0.50:
        whole_body_visible = "partial"
    else:
        whole_body_visible = "no"
    if detection_coverage >= 0.95:
        occlusion = "none"
    elif detection_coverage >= 0.70:
        occlusion = "minor"
    else:
        occlusion = "major"

    return {
        **properties,
        "actual_frames": total_frames,
        "csv_path": csv_path,
        "annotated_path": annotated_path,
        "detection_coverage": detection_coverage,
        "full_body_coverage": full_body_coverage,
        "mean_detection_confidence": float(np.mean(detection_confidences)) if detection_confidences else 0.0,
        "mean_keypoint_confidence": float(np.mean(keypoint_confidences)) if keypoint_confidences else 0.0,
        "max_people": max(person_counts, default=0),
        "whole_body_visible": whole_body_visible,
        "occlusion_level": occlusion,
    }


def next_id(sheet: Any, column: int, prefix: str) -> str:
    highest = 0
    for row in range(5, sheet.max_row + 1):
        value = sheet.cell(row, column).value
        match = re.fullmatch(rf"{re.escape(prefix)}(\d+)", str(value or ""), re.IGNORECASE)
        if match:
            highest = max(highest, int(match.group(1)))
    return f"{prefix}{highest + 1:04d}"


def first_blank_row(sheet: Any, column: int) -> int:
    for row in range(5, max(sheet.max_row, 5) + 1):
        if sheet.cell(row, column).value in {None, ""}:
            return row
    return sheet.max_row + 1


def prepare_row(sheet: Any, row: int, template_row: int = 5) -> None:
    if row <= sheet.max_row:
        return
    for col in range(1, sheet.max_column + 1):
        source = sheet.cell(template_row, col)
        target = sheet.cell(row, col)
        target._style = copy(source._style)
        target.number_format = source.number_format
        target.alignment = copy(source.alignment)
        target.protection = copy(source.protection)
    sheet.row_dimensions[row].height = sheet.row_dimensions[template_row].height
    for validation in sheet.data_validations.dataValidation:
        for cell_range in list(validation.ranges.ranges):
            if cell_range.min_row <= template_row <= cell_range.max_row:
                for col in range(cell_range.min_col, cell_range.max_col + 1):
                    validation.add(sheet.cell(row, col).coordinate)


def relative_path(path: Path) -> str:
    try:
        return str(path.resolve().relative_to(ROOT))
    except ValueError:
        return str(path.resolve())


def update_workbook(
    workbook_path: Path,
    video_path: Path,
    metadata: dict[str, str],
    analysis: dict[str, Any],
    model_path: Path,
) -> tuple[str, str, int, int]:
    if not workbook_path.exists():
        raise FileNotFoundError(f"Dataset workbook not found: {workbook_path}")
    workbook = openpyxl.load_workbook(workbook_path)
    video_sheet = workbook["Videos"]
    clip_sheet = workbook["Movement Clips"]
    video_id = next_id(video_sheet, 1, "V")
    clip_id = next_id(clip_sheet, 1, "C")
    video_row = first_blank_row(video_sheet, 1)
    clip_row = first_blank_row(clip_sheet, 1)
    prepare_row(video_sheet, video_row)
    prepare_row(clip_sheet, clip_row)

    recording_date = datetime.fromtimestamp(video_path.stat().st_mtime).date()
    session_id = f"S-{metadata['participant_id']}-{recording_date.isoformat()}"
    notes = (
        f"Automatic YOLO import. Detection coverage={analysis['detection_coverage']:.1%}; "
        f"side-view body-joint coverage={analysis['full_body_coverage']:.1%}; "
        f"mean keypoint confidence={analysis['mean_keypoint_confidence']:.3f}; "
        f"maximum people detected={analysis['max_people']}."
    )
    video_values = [
        video_id, video_path.name, metadata["participant_id"], session_id, metadata["consent"],
        metadata["exercise_name"], metadata["camera_view"], recording_date,
        round(analysis["duration"], 3), round(analysis["fps"], 3),
        f"{analysis['width']}x{analysis['height']}", analysis["whole_body_visible"], "",
        relative_path(video_path), "unassigned", notes,
    ]
    for column, value in enumerate(video_values, start=1):
        video_sheet.cell(video_row, column, value)

    clip_values = [
        clip_id, video_id, 0.0, round(analysis["duration"], 3), "", "full_rep", "uncertain", "",
        "", "not_applicable", analysis["occlusion_level"], "YOLO auto-import", "needs_review",
        "unassigned", relative_path(analysis["csv_path"]), relative_path(analysis["annotated_path"]),
        model_path.name, "Automatic full-video placeholder. Review timestamps, repetitions, and movement labels.",
    ]
    for column, value in enumerate(clip_values, start=1):
        clip_sheet.cell(clip_row, column, value)

    if workbook.calculation is not None:
        workbook.calculation.fullCalcOnLoad = True
        workbook.calculation.forceFullCalc = True
    workbook_path.parent.mkdir(parents=True, exist_ok=True)
    with tempfile.NamedTemporaryFile(
        prefix=f".{workbook_path.stem}-", suffix=".xlsx", dir=workbook_path.parent, delete=False
    ) as temp_file:
        temp_path = Path(temp_file.name)
    try:
        workbook.save(temp_path)
        os.replace(temp_path, workbook_path)
    finally:
        if temp_path.exists():
            temp_path.unlink()
    return video_id, clip_id, video_row, clip_row


def process_all(args: argparse.Namespace) -> int:
    for directory in [DEFAULT_INBOX, KEYPOINT_DIR, ANNOTATED_DIR]:
        directory.mkdir(parents=True, exist_ok=True)
    videos = discover_videos(args.inputs)
    if not videos:
        print(f"No videos found. Put an MP4, MOV, or other video into: {DEFAULT_INBOX}")
        return 0
    if not args.model.exists():
        raise FileNotFoundError(f"YOLO model not found: {args.model}")
    if not args.workbook.exists():
        raise FileNotFoundError(f"Dataset workbook not found: {args.workbook}")

    manifest = load_manifest()
    model_hash = sha256_file(args.model)[:12]
    device = choose_device(args.device)
    jobs: list[tuple[Path, str, dict[str, str], str]] = []
    for video_path in videos:
        video_hash = sha256_file(video_path)
        job_key = f"{video_hash}:{model_hash}:{args.imgsz}:{args.confidence}"
        if job_key in manifest["processed"]:
            prior = manifest["processed"][job_key]
            print(f"Already processed: {video_path.name} as {prior.get('video_id', 'unknown ID')}")
            continue
        metadata = collect_metadata(video_path, args)
        if metadata["consent"] not in {"yes", "reference_only"}:
            print(f"Skipped {video_path.name}: permission was not confirmed.")
            continue
        output_stem = f"{clean_label(video_path.stem)}_{video_hash[:8]}"
        jobs.append((video_path, output_stem, metadata, job_key))

    if not jobs:
        print("No new permitted videos to process.")
        return 0

    print(f"\nLoading {args.model.name} on {device}...")
    model = YOLO(str(args.model))
    for video_path, output_stem, metadata, job_key in jobs:
        print(f"Processing {video_path.name}...")
        analysis = process_video(
            model, video_path, output_stem, device, args.imgsz, args.confidence, args.keypoint_confidence
        )
        video_id, clip_id, video_row, clip_row = update_workbook(
            args.workbook, video_path, metadata, analysis, args.model
        )
        manifest["processed"][job_key] = {
            "video_id": video_id,
            "clip_id": clip_id,
            "source": relative_path(video_path),
            "model": args.model.name,
            "model_sha256_prefix": model_hash,
            "keypoints_csv": relative_path(analysis["csv_path"]),
            "annotated_video": relative_path(analysis["annotated_path"]),
            "workbook": relative_path(args.workbook),
            "workbook_video_row": video_row,
            "workbook_clip_row": clip_row,
            "processed_at": datetime.now().astimezone().isoformat(timespec="seconds"),
            "ultralytics_version": ultralytics_version,
        }
        save_manifest(manifest)
        print(
            f"Saved {video_id}/{clip_id}: {analysis['actual_frames']} frames, "
            f"{analysis['detection_coverage']:.1%} detection coverage"
        )
        print(f"  Keypoints: {analysis['csv_path']}")
        print(f"  Annotated video: {analysis['annotated_path']}")
        print(f"  Workbook: {args.workbook}")
    return 0


def main() -> int:
    try:
        return process_all(parse_args())
    except KeyboardInterrupt:
        print("\nStopped. The original videos were not changed.", file=sys.stderr)
        return 130
    except Exception as exc:
        print(f"Error: {exc}", file=sys.stderr)
        return 1


if __name__ == "__main__":
    raise SystemExit(main())
